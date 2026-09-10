#!/usr/bin/env python3
"""Derive the pratyāhāra range table from ShivaSutras.xlsx (H4471).

Source: yadisk:Panini/ShivaSutras.xlsx (rights not cleared for
redistribution — the raw file is never committed; fetch it yourself with
``rclone copy yadisk:Panini/ShivaSutras.xlsx data/_raw_local/`` before
running this script).

Schema of the source ``Ranges`` sheet: one row per pratyāhāra (aK, aC, aṬ,
...), column 2 is the pratyāhāra name, column 1 an unexplained numeric tag
(NOT a phoneme count — checked against the sheet's own worked example),
and every row from column 3 onward repeats the *same* 57-symbol alphabet
sequence (letters and the 14 Māheśvara-sūtra anubandha markers alike).
Membership in a pratyāhāra's range is encoded by solid cell fill, not by
cell content — so the range must be read off ``cell.fill.patternType``,
not off which cells contain text.

Output: shiva_sutras_pratyahara_ranges.tsv + .json (42 rows), one row per
pratyāhāra, phonemes in source order including anubandha markers.
"""
import csv
import json
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
SRC = HERE / "_raw_local" / "ShivaSutras.xlsx"
OUT_TSV = HERE / "shiva_sutras_pratyahara_ranges.tsv"
OUT_JSON = HERE / "shiva_sutras_pratyahara_ranges.json"

FIRST_DATA_COL = 3


def main() -> None:
    if not SRC.exists():
        sys.exit(
            f"missing {SRC} — fetch with: "
            f"rclone copy yadisk:Panini/ShivaSutras.xlsx {SRC.parent}/"
        )

    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb["Ranges"]

    rows = []
    for r in range(1, ws.max_row + 1):
        name = ws.cell(row=r, column=2).value
        if not name or not isinstance(name, str):
            continue
        tag = ws.cell(row=r, column=1).value
        span = []
        for c in range(FIRST_DATA_COL, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            if cell.fill and cell.fill.patternType == "solid" and isinstance(cell.value, str):
                span.append(cell.value)
        if not span:
            continue
        rows.append(
            {
                "pratyahara": name,
                "source_tag": tag,
                "span_length": len(span),
                "phonemes_devanagari": span,
            }
        )

    with OUT_TSV.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter="\t")
        w.writerow(["pratyahara", "source_tag", "span_length", "phonemes_devanagari"])
        for row in rows:
            w.writerow(
                [row["pratyahara"], row["source_tag"], row["span_length"], " ".join(row["phonemes_devanagari"])]
            )

    with OUT_JSON.open("w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=2)

    print(f"rows: {len(rows)} -> {OUT_TSV.name}, {OUT_JSON.name}")


if __name__ == "__main__":
    main()
