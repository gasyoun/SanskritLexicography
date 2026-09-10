# -*- coding: utf-8 -*-
r"""Build the Palsule-hubbed multi-source dhātupāṭha coordinate crosswalk (H4478).

SOURCE. MG's own Concordance build, yadisk `Sanskrityatina/09_Palsule/`:
`!Gasuns Concordance/Gasuns-Dhātupāṭha-Concordance (Merge_table_Final).xlsm`,
sheet «Финальная таблица» (== «Final», 7,385 rows × 13 cols). Raw XLSX is
gitignored (H1333 precedent, lands under pwg_ru/eval/); only this DERIVED JSON
is committed.

Schema audit 10-09-2026 (H4478): among the four sibling Final-family sheets —
this one, `22_october.xlsx::Final`, `concordance-PhD.xlsx::october`, and the
`Dhatu_Merge_table_160813` build — this sheet was chosen as authoritative
because it is the only one carrying page and entry № as SEPARATE columns (the
siblings merge them into `стр.-№№№`) and the only one with a populated
«Статистика» sheet (18,457 cells audited by the author).

WHAT A ROW IS. One spreadsheet row, verbatim. A Palsule root entry (1955 artha
index; корень + page + №) side by side with the SAME root's location in five
more dhātupāṭha witnesses: Böhtlingk PWG (том-кол. = volume,column of the PWG
article — the article whose `DHĀTUP. x,y` citations the landed
`dhatup_palsule.json` concordance joins on), Böhtlingk PWK, Whitney Roots
(1885), Mayrhofer EWA (1986-2001), Werba VIA I (1997). Empty cell == that
witness has no entry for this row. Values are copied as-is; nothing is
normalized, matched, or re-ordered — the interpretation layer belongs to the
consumer.

This COMPLEMENTS, never replaces, `src/data/dhatup_palsule.json` (H1333):
that file answers "what does Palsule record for this root spelling"; this file
answers "where does each of the six traditions print this root". Join key
between them: the Palsule root string (same normalization caveats as H1333 —
homonym spellings are NOT disambiguated here either).

Usage: python3 build_dhatup_multisource_crosswalk.py [--xlsx PATH] [--out PATH]
"""
from __future__ import annotations

import argparse
import json
import os
from datetime import date

from openpyxl import load_workbook

DEFAULT_XLSX = os.path.join(
    os.path.dirname(__file__), "..", "pwg_ru", "eval", "palsule_yadisk",
    "Gasuns-Dhātupāṭha-Concordance (Merge_table_Final).xlsm")
DEFAULT_OUT = os.path.join(
    os.path.dirname(__file__), "data", "dhatup_multisource_crosswalk.json")

SHEET = "Финальная таблица"
FIELDS = [
    "palsule_root", "palsule_page", "palsule_no",
    "pwg_volcol", "pwg_no",
    "pwk_volcol", "pwk_no",
    "whitney_root", "whitney_page", "whitney_no",
    "ewa_volcol", "ewa_no",
    "verba_root",
]
HEADER_ROWS = 3  # row1 sources+dates, row2 source years, row3 column names


def _cell(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def build(xlsx_path: str) -> dict:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb[wb.sheetnames[0]]
    rows_out = []
    for i, row in enumerate(ws.iter_rows(min_row=HEADER_ROWS + 1,
                                         max_col=len(FIELDS), values_only=True),
                            start=HEADER_ROWS + 1):
        vals = [_cell(v) for v in row]
        if not any(vals):
            continue  # fully blank grid row
        rec = {"row": i}
        rec.update(dict(zip(FIELDS, vals)))
        rows_out.append(rec)
    wb.close()
    return {
        "_README": __doc__.strip(),
        "_provenance": {
            "handoff": "H4478",
            "built": date.today().isoformat(),
            "source_workbook": os.path.basename(xlsx_path),
            "source_sheet": SHEET,
            "source_yadisk": "Sanskrityatina/09_Palsule/!Gasuns Concordance/",
            "author_of_source": "Dr. Mārcis Gasūns (own Concordance build, 2013-2014)",
            "landed_as": "derived-only (H1333 precedent; raw XLSX gitignored)",
            "sibling_revisions_not_used": [
                "22_october.xlsx::Final (combined стр.-№№№ layout)",
                "22_october.xlsx::c PWK", "concordance-PhD.xlsx::october",
                "Gasuns-Dhātupāṭha-Concordance (Dhatu_Merge_table_160813_Pa-PWG-PWK-Wh-EWA-VIA).xlsm",
            ],
        },
        "counts": {
            "rows": len(rows_out),
            "with_palsule_root": sum(1 for r in rows_out if r["palsule_root"]),
            "per_source_filled": {f: sum(1 for r in rows_out if r[f])
                                  for f in FIELDS},
        },
        "rows": rows_out,
    }


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--xlsx", default=DEFAULT_XLSX)
    ap.add_argument("--out", default=DEFAULT_OUT)
    args = ap.parse_args()
    data = build(args.xlsx)
    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    with open(args.out, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    print(json.dumps(data["counts"], ensure_ascii=False, indent=1))
    print(f"wrote {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
